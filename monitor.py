"""Scrape a product catalog, save a snapshot, and build an Excel report.

The demo target is books.toscrape.com, a public sandbox made for scraping
practice. The same structure works for any shop: swap the parse_* functions.

    python monitor.py                 # full run, writes report + snapshot
    python monitor.py --limit 3       # only the first 3 categories (quick test)
"""

import argparse
import json
import time
from datetime import date
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
UP_FILL = PatternFill("solid", fgColor="F8D7DA")
DOWN_FILL = PatternFill("solid", fgColor="D4EDDA")


def fetch(session, url, retries=3):
    for attempt in range(retries):
        try:
            resp = session.get(url, timeout=20)
            resp.raise_for_status()
            resp.encoding = "utf-8"
            return BeautifulSoup(resp.text, "html.parser")
        except requests.RequestException:
            if attempt == retries - 1:
                raise
            time.sleep(2 * (attempt + 1))


def parse_categories(soup):
    links = soup.select("div.side_categories ul li ul li a")
    return [(a.get_text(strip=True), urljoin(BASE_URL, a["href"])) for a in links]


def parse_listing(soup, page_url, category):
    items = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        price = card.select_one("p.price_color").get_text(strip=True)
        rating_class = card.select_one("p.star-rating")["class"][1]
        items.append({
            "title": link["title"],
            "category": category,
            "price": float(price.replace("£", "").replace("Â", "")),
            "rating": RATINGS.get(rating_class, 0),
            "in_stock": "In stock" in card.select_one("p.availability").get_text(),
            "url": urljoin(page_url, link["href"]),
        })
    next_link = soup.select_one("li.next a")
    next_url = urljoin(page_url, next_link["href"]) if next_link else None
    return items, next_url


def scrape(limit=None, delay=0.3):
    session = requests.Session()
    session.headers["User-Agent"] = "Mozilla/5.0 (catalog-monitor demo)"
    categories = parse_categories(fetch(session, BASE_URL))
    if limit:
        categories = categories[:limit]

    products = []
    for name, url in categories:
        while url:
            items, url = parse_listing(fetch(session, url), url, name)
            products.extend(items)
            time.sleep(delay)
        print(f"  {name}: {sum(p['category'] == name for p in products)} items")
    return products


def compare(products, previous):
    """Return price changes against the previous snapshot, keyed by URL."""
    old = {p["url"]: p for p in previous}
    changes = []
    for p in products:
        before = old.get(p["url"])
        if before and before["price"] != p["price"]:
            changes.append({**p, "old_price": before["price"],
                            "diff": round(p["price"] - before["price"], 2)})
    new_items = [p for p in products if p["url"] not in old]
    return changes, new_items


def style_header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def autosize(ws, max_width=60):
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width, max_width)


def build_report(products, changes, new_items, path):
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    style_header(summary, ["Category", "Products", "Avg price, £",
                           "Min price, £", "Max price, £", "Avg rating"])
    by_cat = {}
    for p in products:
        by_cat.setdefault(p["category"], []).append(p)
    for cat, items in sorted(by_cat.items(), key=lambda kv: -len(kv[1])):
        prices = [i["price"] for i in items]
        summary.append([cat, len(items), round(sum(prices) / len(prices), 2),
                        min(prices), max(prices),
                        round(sum(i["rating"] for i in items) / len(items), 2)])
    for row in summary.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '"£"#,##0.00'
    autosize(summary)

    top = min(len(by_cat), 10)
    chart = BarChart()
    chart.title = "Products per category (top 10)"
    chart.y_axis.title = "Products"
    chart.add_data(Reference(summary, min_col=2, min_row=1, max_row=top + 1), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=2, max_row=top + 1))
    chart.height, chart.width = 9, 18
    chart.legend = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    summary.add_chart(chart, "H2")

    data = wb.create_sheet("All products")
    style_header(data, ["Title", "Category", "Price, £", "Rating", "In stock", "URL"])
    for p in products:
        data.append([p["title"], p["category"], p["price"], p["rating"],
                     "Yes" if p["in_stock"] else "No", p["url"]])
    for cell in data["C"][1:]:
        cell.number_format = '"£"#,##0.00'
    data.auto_filter.ref = data.dimensions
    autosize(data)

    ch = wb.create_sheet("Price changes")
    style_header(ch, ["Title", "Category", "Old price, £", "New price, £", "Change, £", "URL"])
    for c in sorted(changes, key=lambda c: c["diff"], reverse=True):
        ch.append([c["title"], c["category"], c["old_price"], c["price"], c["diff"], c["url"]])
        fill = UP_FILL if c["diff"] > 0 else DOWN_FILL
        for cell in ch[ch.max_row]:
            cell.fill = fill
        for cell in ch[ch.max_row][2:5]:
            cell.number_format = '"£"#,##0.00'
    if not changes:
        ch.append(["No price changes since the previous run."])
    autosize(ch)

    nw = wb.create_sheet("New products")
    style_header(nw, ["Title", "Category", "Price, £", "URL"])
    for p in new_items:
        nw.append([p["title"], p["category"], p["price"], p["url"]])
    autosize(nw)

    wb.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--limit", type=int, help="only scrape the first N categories")
    parser.add_argument("--out", default="output", help="output folder")
    args = parser.parse_args()

    out = Path(args.out)
    out.mkdir(exist_ok=True)
    snapshot_path = out / "last_snapshot.json"
    previous = json.loads(snapshot_path.read_text()) if snapshot_path.exists() else []

    print("Scraping catalog...")
    products = scrape(limit=args.limit)
    changes, new_items = compare(products, previous)
    if not previous:
        new_items = []  # first run: everything is "new", which is not useful

    report = out / f"catalog_report_{date.today():%Y-%m-%d}.xlsx"
    build_report(products, changes, new_items, report)
    snapshot_path.write_text(json.dumps(products, ensure_ascii=False, indent=1))

    print(f"Done: {len(products)} products, {len(changes)} price changes, "
          f"{len(new_items)} new. Report: {report}")


if __name__ == "__main__":
    main()
